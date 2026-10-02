"""
Builds outputs/Flight_Delays_Dashboard.xlsx from the CSVs in data/processed/
(written by export_dashboard_data.py) and outputs/predicted_delays_test_month.csv.

Sheets: Dashboard (KPI cards + 6 charts), Summary (the small tables the charts
plot), and one data sheet per source table. KPI cards are Excel formulas over
the data sheets, so they recalculate if the data sheets are refreshed. Only
aggregates and a labelled sample of flagged flights go into Excel; the full
flight-level data stays in Postgres.
"""
import os
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[2]
PROCESSED_DIR = Path(os.environ.get("DASHBOARD_DATA_DIR", PROJECT_ROOT / "data" / "processed"))
OUTPUTS_DIR = Path(os.environ.get("OUTPUTS_DIR", PROJECT_ROOT / "outputs"))
OUTPUT_PATH = OUTPUTS_DIR / "Flight_Delays_Dashboard.xlsx"
FLAGGED_SAMPLE = 500

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
KPI_FILL = PatternFill(start_color="EFF6FC", end_color="EFF6FC", fill_type="solid")
KPI_LABEL_FONT = Font(size=10, color="44546A", bold=True)
KPI_VALUE_FONT = Font(size=18, bold=True, color="1F4E78")


def write_table(ws, df, name, start_row=1, start_col=1):
    """Writes df as a formatted Excel Table; returns (first_row, last_row)."""
    for j, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(row=start_row, column=j, value=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, value in enumerate(row, start=start_col):
            ws.cell(row=i, column=j, value=None if pd.isna(value) else value)
    last_row = start_row + len(df)
    ref = (f"{get_column_letter(start_col)}{start_row}:"
           f"{get_column_letter(start_col + len(df.columns) - 1)}{last_row}")
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for j, col in enumerate(df.columns, start=start_col):
        ws.column_dimensions[get_column_letter(j)].width = max(12, len(str(col)) + 4)
    return start_row, last_row


def kpi(ws, col, label, formula, number_format):
    letter = get_column_letter(col)
    ws.merge_cells(f"{letter}3:{get_column_letter(col + 1)}3")
    ws.merge_cells(f"{letter}4:{get_column_letter(col + 1)}4")
    label_cell, value_cell = ws[f"{letter}3"], ws[f"{letter}4"]
    label_cell.value, value_cell.value = label, formula
    label_cell.font, value_cell.font = KPI_LABEL_FONT, KPI_VALUE_FONT
    value_cell.number_format = number_format
    for cell in (label_cell, value_cell):
        cell.fill = KPI_FILL
        cell.alignment = Alignment(horizontal="center")


def main():
    read = lambda name, **kw: pd.read_csv(PROCESSED_DIR / f"{name}.csv", **kw)  # noqa: E731
    carriers = read("ontime_by_carrier")
    airports = read("ontime_by_origin")
    hourly = read("ontime_by_hour")
    daymonth = read("ontime_by_dow_month")
    daily = read("ontime_daily", parse_dates=["flight_date"])
    routes = read("top_delayed_routes")
    causes = read("delay_causes_by_month")
    metrics = read("model_metrics")
    by_carrier = read("model_by_carrier")
    importance = read("feature_importance")
    cm = read("confusion_matrix")
    before_after = read("model_before_after")
    flagged_all = pd.read_csv(OUTPUTS_DIR / "predicted_delays_test_month.csv")
    test_month = metrics["test_month"].iloc[0]
    period = f"{daily['flight_date'].min():%b}-{daily['flight_date'].max():%b %Y}"         if daily["flight_date"].min().year == daily["flight_date"].max().year         else f"{daily['flight_date'].min():%b %Y}-{daily['flight_date'].max():%b %Y}"
    flagged = flagged_all.head(FLAGGED_SAMPLE)

    # Chart inputs
    carrier_chart = carriers.sort_values("on_time_pct", ascending=False)[["carrier_name", "on_time_pct"]] \
        .rename(columns={"carrier_name": "Carrier", "on_time_pct": "On-Time %"})
    hour_chart = hourly[["dep_hour", "delay_rate_pct"]].rename(
        columns={"dep_hour": "Scheduled Hour", "delay_rate_pct": "Delay Rate %"})
    hour_chart["Scheduled Hour"] = hour_chart["Scheduled Hour"].map(lambda h: f"{h:02d}:00")
    daily_chart = daily[["flight_date", "on_time_pct"]].rename(
        columns={"flight_date": "Date", "on_time_pct": "On-Time %"})
    causes_chart = (causes.pivot_table(index=["flight_month", "month_name"], columns="cause",
                                       values="delay_minutes", aggfunc="sum")
                    .reset_index().drop(columns="flight_month").rename(columns={"month_name": "Month"}))
    causes_chart = causes_chart[["Month", "Late Aircraft", "Carrier", "National Air System",
                                 "Weather", "Security"]]
    for c in causes_chart.columns[1:]:
        causes_chart[c] = (causes_chart[c] / 1000).round(1)  # thousand minutes
    model_chart = metrics[["model", "auc_pr", "auc_roc", "precision", "recall"]].rename(columns={
        "model": "Model", "auc_pr": "AUC-PR", "auc_roc": "AUC-ROC", "precision": "Precision (top 20%)",
        "recall": "Recall (top 20%)"})
    busiest = airports.nlargest(30, "scheduled_flights")
    airport_chart = (busiest.nlargest(10, "delay_rate_pct")[["airport_code", "delay_rate_pct"]]
                     .rename(columns={"airport_code": "Origin Airport", "delay_rate_pct": "Delay Rate %"}))

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    summary = wb.create_sheet("Summary")
    data_sheets = {
        "Carriers": carriers, "Airports": airports, "Hourly": hourly, "DayMonth": daymonth,
        "Daily": daily, "Routes": routes, "DelayCauses": causes, "ModelMetrics": metrics,
        "ModelByCarrier": by_carrier, "FeatureImportance": importance, "ConfusionMatrix": cm,
        "BeforeAfter": before_after,
    }
    for sheet_name, df in data_sheets.items():
        write_table(wb.create_sheet(sheet_name), df, sheet_name)
    ws_flag = wb.create_sheet("FlaggedFlights")
    ws_flag["A1"] = (f"Top {len(flagged):,} of {len(flagged_all):,} flights the model predicted as delayed "
                     f"in the {test_month} test month, predicted 2 hours before scheduled departure (riskiest 20%, "
                     f"sorted by probability). Full list: outputs/predicted_delays_test_month.csv")
    ws_flag["A1"].font = Font(italic=True, color="44546A")
    write_table(ws_flag, flagged, "FlaggedFlights", start_row=3)
    for sheet in ("Daily",):
        for r in range(2, len(daily) + 2):
            wb[sheet].cell(row=r, column=1).number_format = "yyyy-mm-dd"

    blocks, col = {}, 1
    for name, df in (("CarrierChart", carrier_chart), ("HourChart", hour_chart), ("DailyChart", daily_chart),
                     ("CausesChart", causes_chart), ("ModelChart", model_chart),
                     ("AirportChart", airport_chart)):
        first, last = write_table(summary, df, name, start_col=col)
        blocks[name] = (col, first, last)
        col += len(df.columns) + 1
    dc = blocks["DailyChart"][0]
    for r in range(2, len(daily_chart) + 2):
        summary.cell(row=r, column=dc).number_format = "dd mmm"

    # Dashboard: title, KPI cards (formulas over the data sheets), charts
    dash["A1"] = f"US Airline Flight Delays Dashboard (BTS On-Time Performance, {period})"
    dash["A1"].font = Font(size=20, bold=True, color="1F4E78")
    dash["A2"] = ("Operations KPIs over all scheduled flights. Model KPIs: LightGBM predicting a 15+ min arrival delay "
                  "2 hours before scheduled departure (inbound aircraft, airport congestion, weather, schedule), "
                  f"fit on earlier months and scored once on the unseen {test_month} flights; precision/recall: each "
                  "model flags its riskiest 20% (or 10%).")
    dash["A2"].font = Font(italic=True, color="44546A")
    kpi(dash, 1, "Total Scheduled Flights", "=SUM(Carriers[scheduled_flights])", "#,##0")
    kpi(dash, 3, "On-Time Arrival %", "=1-SUM(Carriers[delayed_flights])/SUM(Carriers[completed_flights])", "0.0%")
    kpi(dash, 5, "Cancellation Rate", "=SUM(Carriers[cancelled_flights])/SUM(Carriers[scheduled_flights])", "0.00%")
    kpi(dash, 7, "Model AUC-PR (headline)",
        '=INDEX(ModelMetrics[auc_pr],MATCH("LightGBM",ModelMetrics[model],0))', "0.000")
    kpi(dash, 9, "Precision (riskiest 20%)",
        '=INDEX(ModelMetrics[precision],MATCH("LightGBM",ModelMetrics[model],0))', "0.0%")
    kpi(dash, 11, "Recall (riskiest 20%)",
        '=INDEX(ModelMetrics[recall],MATCH("LightGBM",ModelMetrics[model],0))', "0.0%")
    kpi(dash, 13, "Precision (riskiest 10%)",
        '=INDEX(ModelMetrics[precision_at_10pct],MATCH("LightGBM",ModelMetrics[model],0))', "0.0%")
    kpi(dash, 15, "Recall (riskiest 10%)",
        '=INDEX(ModelMetrics[recall_at_10pct],MATCH("LightGBM",ModelMetrics[model],0))', "0.0%")
    lookup = 'TEXT(INDEX(ModelMetrics[auc_pr],MATCH("{}",ModelMetrics[model],0)),"0.000")'
    dash["A5"] = ('="AUC-PR baselines: historical carrier+route rate "&'
                  + lookup.format("Historical rate rule (carrier+route)")
                  + '&", inbound-aircraft delay rule "&' + lookup.format("Inbound-delay rule")
                  + '&", logistic regression "&' + lookup.format("Logistic regression")
                  + '&". A random guess scores the test-month delay rate: "&TEXT(SUMIFS(ConfusionMatrix[flights],'
                  'ConfusionMatrix[actual],"Delayed")/SUM(ConfusionMatrix[flights]),"0.000")')
    dash["A5"].font = Font(italic=True, color="44546A")
    for c in range(1, 17):
        dash.column_dimensions[get_column_letter(c)].width = 14

    def ref(block, first_col, last_col):
        c0, first, last = blocks[block]
        return Reference(summary, min_col=c0 + first_col, max_col=c0 + last_col, min_row=first, max_row=last)

    def cats(block):
        c0, first, last = blocks[block]
        return Reference(summary, min_col=c0, min_row=first + 1, max_row=last)

    carrier_bar = BarChart()
    carrier_bar.type = "bar"
    carrier_bar.title = "On-Time Arrival % by Carrier"
    carrier_bar.add_data(ref("CarrierChart", 1, 1), titles_from_data=True)
    carrier_bar.set_categories(cats("CarrierChart"))
    carrier_bar.legend = None

    hour_bar = BarChart()
    hour_bar.type = "col"
    hour_bar.title = "Delay Rate % by Scheduled Departure Hour"
    hour_bar.add_data(ref("HourChart", 1, 1), titles_from_data=True)
    hour_bar.set_categories(cats("HourChart"))
    hour_bar.legend = None

    daily_line = LineChart()
    daily_line.title = "Daily On-Time Arrival %"
    daily_line.add_data(ref("DailyChart", 1, 1), titles_from_data=True)
    daily_line.set_categories(cats("DailyChart"))
    daily_line.x_axis.number_format = "dd mmm"
    daily_line.legend = None

    causes_bar = BarChart()
    causes_bar.type, causes_bar.grouping, causes_bar.overlap = "col", "stacked", 100
    causes_bar.title = "Delay Minutes by Cause (thousands, 15+ min late flights)"
    causes_bar.add_data(ref("CausesChart", 1, 5), titles_from_data=True)
    causes_bar.set_categories(cats("CausesChart"))

    model_bar = BarChart()
    model_bar.type, model_bar.grouping = "bar", "clustered"
    model_bar.title = f"Model Comparison, {test_month} Test Flights, 2 h Before Departure (riskiest 20%)"
    model_bar.add_data(ref("ModelChart", 1, 4), titles_from_data=True)
    model_bar.set_categories(cats("ModelChart"))

    airport_bar = BarChart()
    airport_bar.type = "bar"
    airport_bar.title = "Most-Delayed of the 30 Busiest Origin Airports (Delay Rate %)"
    airport_bar.add_data(ref("AirportChart", 1, 1), titles_from_data=True)
    airport_bar.set_categories(cats("AirportChart"))
    airport_bar.legend = None

    # On horizontal bar charts x_axis is the category axis; maxMin lists rows top-down,
    # and crosses='max' keeps the value axis at the bottom instead of under the title.
    for bar in (carrier_bar, model_bar, airport_bar):
        bar.x_axis.scaling.orientation = "maxMin"
        bar.y_axis.crosses = "max"
    carrier_bar.y_axis.scaling.min = 50
    carrier_bar.y_axis.scaling.max = 100

    layout = ((carrier_bar, "A7"), (hour_bar, "G7"), (daily_line, "A25"), (causes_bar, "G25"),
              (model_bar, "A43"), (airport_bar, "G43"))
    for chart, anchor in layout:
        # openpyxl leaves these unset, and current Excel then hides both axes and
        # colours every point differently.
        chart.varyColors = False
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        for series in chart.series:
            series.smooth = False
        chart.width, chart.height = 17, 8.5
        dash.add_chart(chart, anchor)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}: sheets={wb.sheetnames}, charts on Dashboard={len(dash._charts)}")


if __name__ == "__main__":
    main()
